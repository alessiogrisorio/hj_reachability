from __future__ import annotations

import json
import os
import sys
from dataclasses import dataclass
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# ---- Configurazione ---------------------------------------------------------

METRICS = [
    "euclidean",
    "dce",
    "ttc",
    "eggert",
    "time_eggert",
    "rss",
    "sff",
]

# Valori di esempio: modificarli in base alla scala numerica di ogni metrica.
# Selezione locale: |BRT| < EPSILON[metric].
EPSILON = {
    "euclidean": 0.5,
    "dce": 0.5,
    "ttc": 0.5,
    "eggert": 0.5,
    "time_eggert": 0.5,
    "rss": 0.5,
    "sff": 0.5,
}

# Matrice N x 2: tutte le 21 coppie possibili tra le 7 metriche.
# Eliminare o commentare le righe per limitare i confronti.
IOU_PAIRS = [
    ("euclidean", "dce"),
    ("euclidean", "ttc"),
    ("euclidean", "eggert"),
    ("euclidean", "time_eggert"),
    ("euclidean", "rss"),
    ("euclidean", "sff"),
    ("dce", "ttc"),
    ("dce", "eggert"),
    ("dce", "time_eggert"),
    ("dce", "rss"),
    ("dce", "sff"),
    ("ttc", "eggert"),
    ("ttc", "time_eggert"),
    ("ttc", "rss"),
    ("ttc", "sff"),
    ("eggert", "time_eggert"),
    ("eggert", "rss"),
    ("eggert", "sff"),
    ("time_eggert", "rss"),
    ("time_eggert", "sff"),
    ("rss", "sff"),
]

# Parametri di rappresentazione (non alterano le statistiche nell'Excel).
# Poche colonne consentono di leggere il minimo e il massimo sotto ogni barra.
N_BINS = 10
HISTOGRAM_BAR_WIDTH = 0.74  # < 1: lascia uno spazio visibile tra le colonne
HISTOGRAM_DPI = 200
HISTOGRAM_X_PERCENTILES = (0.5, 100.0)
HISTOGRAM_PERCENTAGE = True
HISTOGRAM_SHOW_P3 = True
HISTOGRAM_SHOW_UNSAFE = True
HISTOGRAM_HANDLE_SATURATION = False

# Per queste metriche si cercano saturazioni ai limiti inferiore e superiore
# di V0: si omettono dalle sole barre se la frequenza supera la soglia.
# Non vengono mai escluse dalle statistiche e dalle percentuali unsafe.
HISTOGRAM_SATURATION_METRICS = ("ttc", "eggert", "time_eggert", "rss")
HISTOGRAM_SATURATION_ATOL = 1e-6
HISTOGRAM_SATURATION_RTOL = 1e-7
HISTOGRAM_SATURATION_MIN_FRACTION = 0.02

GRADIENT_BATCH_LINES = 2048

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
NPZ_DIR = PROJECT_ROOT / "results" / "brt"
OUTPUT_PATH = SCRIPT_DIR / "statistiche.xlsx"
PLOT_DIR = SCRIPT_DIR / "statistiche_plots"

STATE_NAMES = ("x_rel", "y_rel", "theta_rel", "v_H", "delta_E", "v_E")

VALUE_ROWS_BASE = (
    "V0 min",
    "V0 max",
    "BRT min",
    "BRT max",
    "V0 > 0 [%]",
    "V0 <= 0 [%]",
    "BRT > 0 [%]",
    "BRT <= 0 [%]",
    "V0 > 0 -> BRT <= 0 [%]",
    "Delta V media",
    "Delta V deviazione standard",
)
VALUE_ROWS = (
    *VALUE_ROWS_BASE,
    "Rapporto unsafe BRT/V0",
    "BRT percentile 3 (97% >= soglia)",
    *(
        f"{extreme} state: {state}"
        for extreme in ("V0 min", "V0 max", "BRT min", "BRT max")
        for state in STATE_NAMES
    ),
)
GRADIENT_ROWS = tuple(
    label
    for state in STATE_NAMES
    for label in (
        f"dV0/d{state} min",
        f"dV0/d{state} max",
        f"dV0/d{state} std",
        f"dBRT/d{state} min",
        f"dBRT/d{state} max",
        f"dBRT/d{state} std",
        f"dDelta/d{state} std",
        f"dDelta/d{state} mean abs",
    )
)
ZERO_ROWS = (
    "Epsilon |BRT|",
    "N punti vicino a BRT=0",
    "Punti vicino a BRT=0 [%]",
    *(
        label
        for state in STATE_NAMES
        for label in (f"dDelta/d{state} std", f"dDelta/d{state} mean abs")
    ),
)


# ---- Statistiche ------------------------------------------------------------

@dataclass
class RunningStats:
    count: int = 0
    mean: float = 0.0
    m2: float = 0.0
    minimum: float = float("inf")
    maximum: float = float("-inf")
    abs_sum: float = 0.0

    def update(self, data: np.ndarray) -> None:
        clean = np.asarray(data)[np.isfinite(data)]
        if clean.size == 0:
            return
        clean = clean.astype(np.float64, copy=False)
        n = clean.size
        chunk_mean = float(np.mean(clean))
        chunk_m2 = float(np.sum((clean - chunk_mean) ** 2))
        if self.count:
            d = chunk_mean - self.mean
            total = self.count + n
            self.m2 += chunk_m2 + d * d * self.count * n / total
            self.mean += d * n / total
            self.count = total
        else:
            self.count = n
            self.mean = chunk_mean
            self.m2 = chunk_m2
        self.minimum = min(self.minimum, float(np.min(clean)))
        self.maximum = max(self.maximum, float(np.max(clean)))
        self.abs_sum += float(np.sum(np.abs(clean)))

    @property
    def std(self) -> float | None:
        return float(np.sqrt(max(self.m2, 0.0) / self.count)) if self.count else None

    @property
    def mean_abs(self) -> float | None:
        return self.abs_sum / self.count if self.count else None

    @property
    def min_or_none(self) -> float | None:
        return self.minimum if self.count else None

    @property
    def max_or_none(self) -> float | None:
        return self.maximum if self.count else None


def finite_limits(values: np.ndarray, metric: str, description: str):
    finite = np.isfinite(values)
    n_invalid = values.size - np.count_nonzero(finite)
    if n_invalid:
        print(f"  Attenzione: {description} contiene {n_invalid} valori non finiti ({metric}).")
    if n_invalid == values.size:
        return None, None
    return (
        float(np.min(values, where=finite, initial=np.inf)),
        float(np.max(values, where=finite, initial=-np.inf)),
    )


def extrema_states(values: np.ndarray, coords: tuple[np.ndarray, ...]):
    finite = np.isfinite(values)
    if not np.any(finite):
        return [None] * 6, [None] * 6
    idx_min = np.unravel_index(np.argmin(np.where(finite, values, np.inf)), values.shape)
    idx_max = np.unravel_index(np.argmax(np.where(finite, values, -np.inf)), values.shape)
    return (
        [float(coords[i][idx_min[i]]) for i in range(6)],
        [float(coords[i][idx_max[i]]) for i in range(6)],
    )


def value_statistics(v0, brt, coords, metric):
    if v0.shape != brt.shape or v0.ndim != 6:
        raise ValueError(f"{metric}: V0 e BRT devono avere la stessa forma 6D.")

    n_total = v0.size
    v0_min, v0_max = finite_limits(v0, metric, "V0")
    brt_min, brt_max = finite_limits(brt, metric, "BRT")

    finite_v0 = np.isfinite(v0)
    finite_brt = np.isfinite(brt)
    v0_positive = finite_v0 & (v0 > 0)
    v0_nonpositive = finite_v0 & (v0 <= 0)
    brt_positive = finite_brt & (brt > 0)
    brt_nonpositive = finite_brt & (brt <= 0)
    count_v0_unsafe = int(np.count_nonzero(v0_nonpositive))
    count_brt_unsafe = int(np.count_nonzero(brt_nonpositive))
    pct = lambda n: 100.0 * n / n_total

    delta = brt - v0
    finite_delta = finite_v0 & finite_brt & np.isfinite(delta)
    if np.any(finite_delta):
        valid_delta = delta[finite_delta]
        mean_delta = float(np.mean(valid_delta, dtype=np.float64))
        std_delta = float(np.std(valid_delta, dtype=np.float64, ddof=0))
    else:
        mean_delta = None
        std_delta = None

    valid_brt = brt[finite_brt]
    # P3: circa il 97% dei valori e' >= a questa soglia.
    percentile_3 = float(np.percentile(valid_brt, 3)) if valid_brt.size else None
    v0_min_state, v0_max_state = extrema_states(v0, coords)
    brt_min_state, brt_max_state = extrema_states(brt, coords)

    return (
        v0_min, v0_max, brt_min, brt_max,
        pct(np.count_nonzero(v0_positive)),
        pct(count_v0_unsafe),
        pct(np.count_nonzero(brt_positive)),
        pct(count_brt_unsafe),
        pct(np.count_nonzero(v0_positive & brt_nonpositive)),
        mean_delta, std_delta,
        (count_brt_unsafe / count_v0_unsafe if count_v0_unsafe else None),
        percentile_3,
        *v0_min_state, *v0_max_state, *brt_min_state, *brt_max_state,
    )


# ---- Gradienti --------------------------------------------------------------

def read_solver_accuracy(archive) -> str:
    if "metadata_json" not in archive:
        raise ValueError("metadata_json assente: impossibile identificare lo schema dei gradienti salvati.")
    metadata = json.loads(str(archive["metadata_json"].item()))
    accuracy = metadata.get("solver", {}).get("accuracy")
    if accuracy not in {"low", "medium", "high", "very_high"}:
        raise ValueError(f"Accuratezza del solver non riconosciuta: {accuracy!r}")
    return accuracy


def gradient_statistics(v0, brt, gradients, coords, periodic_dims, accuracy, epsilon, metric):
    """Calcola dV0 con lo stesso schema del BRT salvato, una dimensione per volta.

    I gradienti V0 non vengono conservati come array 6D completo.
    """
    if gradients.shape != (*v0.shape, 6):
        raise ValueError(
            f"{metric}: gradients.shape={gradients.shape}; atteso {(*v0.shape, 6)}."
        )

    # Gli stessi operatori usati in hj.Grid.grad_values: media tra WENO/ENO left e right.
    os.environ.setdefault("JAX_PLATFORMS", "cpu")
    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    import jax
    import jax.numpy as jnp
    from hj_reachability import boundary_conditions
    from hj_reachability.finite_differences import upwind_first

    schemes = {
        "low": upwind_first.first_order,
        "medium": upwind_first.ENO2,
        "high": upwind_first.WENO3,
        "very_high": upwind_first.WENO5,
    }
    scheme = schemes[accuracy]

    results = []
    local_results = [float(epsilon)]
    n_local = int(np.count_nonzero(np.isfinite(brt) & (np.abs(brt) < epsilon)))
    local_results.extend((n_local, 100.0 * n_local / brt.size))
    if n_local == 0:
        print(f"  Attenzione: {metric}, nessun punto con |BRT| < {epsilon:g}.")

    for axis, name in enumerate(STATE_NAMES):
        vector = coords[axis]
        if len(vector) < 2:
            raise ValueError(f"{metric}: servono almeno 2 nodi per {name}.")
        spacing = float(vector[1] - vector[0])
        periodic = axis in periodic_dims
        bc = (boundary_conditions.periodic if periodic
              else boundary_conditions.extrapolate_away_from_zero)

        def gradient_line(line):
            left, right = scheme(line, spacing, bc)
            return (left + right) / 2.0

        gradient_batch = jax.jit(jax.vmap(gradient_line, in_axes=1, out_axes=1))

        n_axis = len(vector)
        v0_lines = np.moveaxis(v0, axis, 0).reshape(n_axis, -1)
        brt_lines = np.moveaxis(brt, axis, 0).reshape(n_axis, -1)
        brt_gradient_lines = np.moveaxis(gradients[..., axis], axis, 0).reshape(n_axis, -1)

        initial_stats = RunningStats()
        brt_stats = RunningStats()
        delta_stats = RunningStats()
        near_zero_stats = RunningStats()

        n_lines = v0_lines.shape[1]
        for start in range(0, n_lines, GRADIENT_BATCH_LINES):
            stop = min(start + GRADIENT_BATCH_LINES, n_lines)
            initial_part = np.asarray(gradient_batch(jnp.asarray(v0_lines[:, start:stop])))
            brt_part = brt_gradient_lines[:, start:stop]
            delta_part = brt_part.astype(np.float64) - initial_part.astype(np.float64)
            valid_pairs = np.isfinite(brt_part) & np.isfinite(initial_part)
            near_zero = (valid_pairs & np.isfinite(brt_lines[:, start:stop]) &
                         (np.abs(brt_lines[:, start:stop]) < epsilon))

            initial_stats.update(initial_part)
            brt_stats.update(brt_part)
            delta_stats.update(delta_part[valid_pairs])
            near_zero_stats.update(delta_part[near_zero])

        results.extend((
            initial_stats.min_or_none, initial_stats.max_or_none, initial_stats.std,
            brt_stats.min_or_none, brt_stats.max_or_none, brt_stats.std,
            delta_stats.std, delta_stats.mean_abs,
        ))
        local_results.extend((near_zero_stats.std, near_zero_stats.mean_abs))
        if initial_stats.count != v0.size or brt_stats.count != brt.size:
            print(f"  Attenzione: alcuni gradienti non finiti per {metric}, {name}.")
        if near_zero_stats.count != n_local:
            print(f"  Attenzione: {metric}, {name}: gradienti non finiti nella regione locale.")
        print(f"  {metric}: gradienti d/d{name} completati")

        del gradient_batch, v0_lines, brt_lines, brt_gradient_lines

    return tuple(results), tuple(local_results)


# ---- Visualizzazioni e confronto tra metriche -------------------------------

def histogram_limits(values: tuple[np.ndarray, np.ndarray]) -> tuple[float, float]:
    """Limiti comuni X robusti: percentili calcolati separatamente per V0/BRT."""
    lower, upper = HISTOGRAM_X_PERCENTILES
    available = [x for x in values if x.size]
    if not available:
        return -1.0, 1.0

    lo = min(float(np.percentile(x, lower)) for x in available)
    hi = max(float(np.percentile(x, upper)) for x in available)

    if not (np.isfinite(lo) and np.isfinite(hi)):
        raise ValueError("Impossibile determinare i limiti degli istogrammi.")
    if lo == hi:
        scale = max(abs(lo) * 0.05, 0.05)
        lo, hi = lo - scale, hi + scale

    # Un piccolo margine evita che i valori ai limiti coincidano con la cornice.
    margin = (hi - lo) * 0.025
    return lo - margin, hi + margin


def histogram_data(values: np.ndarray, caps: dict[str, float]):
    """Restituisce dati finiti, dati non saturati e quote dei singoli picchi."""
    finite = values[np.isfinite(values)]
    if finite.size == 0:
        return finite, finite, {}
    saturated = np.zeros(finite.shape, dtype=bool)
    fractions = {}
    for label, cap in caps.items():
        match = np.isclose(
            finite, cap,
            atol=HISTOGRAM_SATURATION_ATOL,
            rtol=HISTOGRAM_SATURATION_RTOL,
        )
        fractions[label] = 100.0 * np.count_nonzero(match) / finite.size
        saturated |= match
    return finite, finite[~saturated], fractions


def plot_histograms(v0, brt, metric, percentile_3):
    """Istogrammi a colonne separate, con minimo e massimo sotto ogni colonna."""
    finite_v0 = v0[np.isfinite(v0)]
    finite_brt = brt[np.isfinite(brt)]
    if not finite_v0.size or not finite_brt.size:
        print(f"  Istogramma saltato: {metric}, nessun dato finito.")
        return

    caps = {}
    if HISTOGRAM_HANDLE_SATURATION and metric in HISTOGRAM_SATURATION_METRICS:
        for label, candidate in (("min", float(np.min(finite_v0))),
                                 ("max", float(np.max(finite_v0)))):
            if any(np.isclose(candidate, other,
                              atol=HISTOGRAM_SATURATION_ATOL,
                              rtol=HISTOGRAM_SATURATION_RTOL)
                   for other in caps.values()):
                continue
            fraction = np.count_nonzero(np.isclose(
                finite_v0, candidate,
                atol=HISTOGRAM_SATURATION_ATOL,
                rtol=HISTOGRAM_SATURATION_RTOL,
            )) / finite_v0.size
            if fraction >= HISTOGRAM_SATURATION_MIN_FRACTION:
                caps[label] = candidate
        if not caps:
            print(f"  {metric}: nessun picco di saturazione identificato in V0.")

    a, a_visible, sat_a = histogram_data(v0, caps)
    b, b_visible, sat_b = histogram_data(brt, caps)
    lo, hi = histogram_limits((a_visible, b_visible))

    # Gli intervalli sono definiti una sola volta, condivisi da V0 e BRT.
    edges = np.linspace(lo, hi, N_BINS + 1)
    counts_v0, _ = np.histogram(a_visible, bins=edges)
    counts_brt, _ = np.histogram(b_visible, bins=edges)

    outside_a = 100.0 * (a_visible.size - int(counts_v0.sum())) / a.size
    outside_b = 100.0 * (b_visible.size - int(counts_brt.sum())) / b.size
    unsafe_a = 100.0 * np.count_nonzero(a <= 0) / a.size
    unsafe_b = 100.0 * np.count_nonzero(b <= 0) / b.size

    if HISTOGRAM_PERCENTAGE:
        hist_a = counts_v0 * (100.0 / a.size)
        hist_b = counts_brt * (100.0 / b.size)
        y_label = "Frequenza [% dei nodi]"
    else:
        hist_a = counts_v0
        hist_b = counts_brt
        y_label = "Frequenza [nodi]"

    ymax = 1.12 * max(float(np.max(hist_a)), float(np.max(hist_b)), 1e-10)
    positions = np.arange(N_BINS)

    # Ogni etichetta riporta esplicitamente i due estremi del relativo bin.
    # Le cifre decimali si adattano all'ampiezza per distinguere bin vicini.
    step = float(edges[1] - edges[0])
    decimals = min(6, max(2, int(np.ceil(-np.log10(step))) + 1))
    labels = [f"{edges[i]:.{decimals}f}\n{edges[i + 1]:.{decimals}f}"
              for i in range(N_BINS)]

    figure_width = min(32.0, max(18.0, 0.87 * N_BINS))
    fig, axes = plt.subplots(2, 1, figsize=(figure_width, 11.8),
                             sharex=True, sharey=True)
    specs = (
        ("V0", hist_a, a, sat_a, outside_a, unsafe_a, "tab:blue"),
        ("BRT (-3 s)", hist_b, b, sat_b, outside_b, unsafe_b, "tab:orange"),
    )

    def value_position(value):
        """Proietta un valore fisico sull'asse a colonne equispaziate."""
        return N_BINS * (value - lo) / (hi - lo) - 0.5

    for ax, (title, hist, data, saturation, outside, unsafe, color) in zip(axes, specs):
        ax.bar(
            positions, hist, width=HISTOGRAM_BAR_WIDTH,
            align="center", color=color, alpha=0.87,
            edgecolor="white", linewidth=0.9,
            zorder=3,
        )
        # V=0 e P3 restano posizionati in base al loro valore numerico reale.
        if lo <= 0.0 <= hi:
            ax.axvline(value_position(0.0), color="black", ls="--", lw=1.3, zorder=5)
        if HISTOGRAM_SHOW_UNSAFE:
            title += f"  |  Unsafe: {unsafe:.2f}%"
        for label, cap in caps.items():
            title += f"  |  Sat. {label} ({cap:.4g}): {saturation[label]:.2f}%"
        if outside > 0.005:
            title += f"  |  Fuori zoom: {outside:.2f}%"
        ax.set_title(title, loc="left", fontsize=11, pad=11)
        ax.set_ylabel(y_label)
        ax.set_xlim(-0.65, N_BINS - 0.35)
        ax.set_ylim(0, ymax)
        ax.set_xticks(positions)
        ax.set_xticklabels(labels, fontsize=7.8, linespacing=1.3)
        # Anche l'istogramma superiore mostra il range alla base di ogni barra.
        ax.tick_params(axis="x", labelbottom=True, length=0, pad=8)
        ax.set_xlabel("Intervallo di V per colonna: minimo (sopra) / massimo (sotto)",
                      labelpad=9, fontsize=9.5)
        ax.grid(axis="y", color="0.88", linewidth=0.7, zorder=0)
        ax.set_axisbelow(True)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        ax.text(
            0.995, 0.97,
            f"Min reale: {np.min(data):.4g}   |   Max reale: {np.max(data):.4g}",
            ha="right", va="top", transform=ax.transAxes,
            fontsize=9, color="0.3",
            bbox=dict(facecolor="white", edgecolor="none", alpha=0.85, pad=2.5),
        )

    if HISTOGRAM_SHOW_P3 and percentile_3 is not None and lo <= percentile_3 <= hi:
        axes[1].axvline(value_position(percentile_3), color="tab:red",
                        ls=":", lw=2, zorder=5)

    from matplotlib.lines import Line2D
    legend = [Line2D([0], [0], color="black", ls="--", lw=1.3, label="V = 0")]
    if HISTOGRAM_SHOW_P3 and percentile_3 is not None:
        suffix = "" if lo <= percentile_3 <= hi else " (fuori zoom)"
        legend.append(Line2D([0], [0], color="tab:red", ls=":", lw=2,
                             label=f"P3 BRT = {percentile_3:.4g} (97% >=){suffix}"))

    fig.suptitle(f"Distribuzione delle Value Functions — {metric}", y=0.985, fontsize=15)
    fig.legend(handles=legend, loc="upper center", bbox_to_anchor=(0.5, 0.957),
               ncol=len(legend), frameon=False, fontsize=10)
    footer = (f"Intervalli: {N_BINS}; estremi riportati sotto ogni colonna. "
              f"Zoom: P{HISTOGRAM_X_PERCENTILES[0]:g}–P{HISTOGRAM_X_PERCENTILES[1]:g} "
              "dei valori non saturati. Percentuali riferite a tutti i nodi finiti.")
    if caps:
        footer += " I picchi saturati sono indicati nei titoli e omessi dalle colonne."
    fig.text(0.5, 0.035, footer, ha="center", va="center", fontsize=9, color="0.35")

    # Spazio tra i pannelli sufficiente per mostrare due righe di numeri per bin.
    fig.subplots_adjust(top=0.875, bottom=0.155, left=0.06, right=0.985, hspace=0.64)
    fig.savefig(PLOT_DIR / f"{metric}_histograms.png", dpi=HISTOGRAM_DPI)
    plt.close(fig)


# Precalcolo del numero di bit 1 per ciascun possibile byte.
BIT_COUNT = np.unpackbits(np.arange(256, dtype=np.uint8)[:, None], axis=1).sum(axis=1)


def count_common_bits(a: np.ndarray, b: np.ndarray) -> int:
    return int(np.sum(BIT_COUNT[np.bitwise_and(a, b)], dtype=np.int64))


def compute_iou(packed_masks: dict, pair_list: list[tuple[str, str]]):
    metrics = [name for name in METRICS if any(name in pair for pair in pair_list)]
    values = np.full((len(metrics), len(metrics)), np.nan)
    for i in range(len(metrics)):
        values[i, i] = 1.0

    index = {name: i for i, name in enumerate(metrics)}
    results = []
    for first, second in pair_list:
        mask_a, n_a, coords_a, periodic_a = packed_masks[first]
        mask_b, n_b, coords_b, periodic_b = packed_masks[second]
        same_grid = (periodic_a == periodic_b and all(
            va.shape == vb.shape and np.allclose(va, vb, rtol=0, atol=1e-6)
            for va, vb in zip(coords_a, coords_b)
        ))
        if not same_grid:
            raise ValueError(f"IoU tra {first} e {second}: le griglie non coincidono.")
        intersection = count_common_bits(mask_a, mask_b)
        union = n_a + n_b - intersection
        iou = intersection / union if union else 1.0
        values[index[first], index[second]] = iou
        values[index[second], index[first]] = iou
        results.append((first, second, iou))
    return metrics, values, results


def plot_iou_matrix(metrics, values):
    if not metrics:
        return
    fig, ax = plt.subplots(figsize=(max(6, len(metrics) * 1.2), max(5, len(metrics))))
    matrix = np.ma.masked_invalid(values)
    im = ax.imshow(matrix, vmin=0, vmax=1, cmap="Blues")
    ax.set_xticks(range(len(metrics)), labels=metrics, rotation=40, ha="right")
    ax.set_yticks(range(len(metrics)), labels=metrics)
    ax.set_title("Intersection over Union - BRT unsafe sets")
    for i in range(len(metrics)):
        for j in range(len(metrics)):
            if np.isfinite(values[i, j]):
                val = values[i, j]
                ax.text(j, i, f"{val:.3f}", ha="center", va="center",
                        color="white" if val > 0.6 else "black", fontsize=9)
    fig.colorbar(im, ax=ax, label="IoU")
    fig.tight_layout()
    fig.savefig(PLOT_DIR / "iou_matrix.png", dpi=HISTOGRAM_DPI)
    plt.close(fig)


# ---- Excel ------------------------------------------------------------------

def write_sheet(sheet, labels, data):
    sheet.append(["Statistica", *METRICS])
    for i, label in enumerate(labels):
        sheet.append([label, *(data[metric][i] for metric in METRICS)])

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 42
    sheet.row_dimensions[1].height = 24

    for col in range(1, len(METRICS) + 2):
        cell = sheet.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if col > 1:
            sheet.column_dimensions[get_column_letter(col)].width = 19

    for row in range(2, len(labels) + 2):
        sheet.cell(row=row, column=1).font = Font(bold=True)
        is_percent = labels[row - 2].endswith("[%]")
        is_integer = labels[row - 2] == "N punti vicino a BRT=0"
        for col in range(2, len(METRICS) + 2):
            cell = sheet.cell(row=row, column=col)
            cell.number_format = '0.00"%"' if is_percent else ("0" if is_integer else "0.000000")
            cell.alignment = Alignment(horizontal="right")


def write_iou_sheet(sheet, metrics, matrix, pairs):
    sheet.append(["BRT unsafe IoU", *metrics])
    for i, name in enumerate(metrics):
        sheet.append([name, *(None if np.isnan(v) else float(v) for v in matrix[i])])
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 23
    for col in range(1, len(metrics) + 2):
        sheet.column_dimensions[get_column_letter(col)].width = 18
        cell = sheet.cell(1, col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    for row in sheet.iter_rows(min_row=2, min_col=2, max_row=1 + len(metrics)):
        for cell in row:
            cell.number_format = "0.0000"
            cell.alignment = Alignment(horizontal="center")
    start_row = len(metrics) + 4
    sheet.cell(start_row, 1, "Coppia 1").font = Font(bold=True)
    sheet.cell(start_row, 2, "Coppia 2").font = Font(bold=True)
    sheet.cell(start_row, 3, "IoU").font = Font(bold=True)
    for rownum, (first, second, value) in enumerate(pairs, start_row + 1):
        sheet.cell(rownum, 1, first)
        sheet.cell(rownum, 2, second)
        sheet.cell(rownum, 3, value).number_format = "0.0000"
    if not metrics:
        sheet.cell(2, 1, "Nessuna coppia configurata in IOU_PAIRS.")


# ---- Esecuzione -------------------------------------------------------------

def validate_configuration():
    if not METRICS or len(METRICS) != len(set(METRICS)):
        raise ValueError("METRICS deve contenere metriche senza duplicati.")
    if N_BINS <= 0 or GRADIENT_BATCH_LINES <= 0 or HISTOGRAM_DPI <= 0:
        raise ValueError("N_BINS, GRADIENT_BATCH_LINES e HISTOGRAM_DPI devono essere positivi.")
    if not (0.0 < HISTOGRAM_BAR_WIDTH < 1.0):
        raise ValueError("HISTOGRAM_BAR_WIDTH deve essere compreso tra 0 e 1 (esclusi).")
    lower, upper = HISTOGRAM_X_PERCENTILES
    if not (0 <= lower < upper <= 100):
        raise ValueError("HISTOGRAM_X_PERCENTILES deve contenere due percentili crescenti in [0,100].")
    if any(name not in METRICS for name in HISTOGRAM_SATURATION_METRICS):
        raise ValueError("HISTOGRAM_SATURATION_METRICS contiene metriche non presenti in METRICS.")
    if not (0 <= HISTOGRAM_SATURATION_MIN_FRACTION <= 1):
        raise ValueError("HISTOGRAM_SATURATION_MIN_FRACTION deve essere compreso tra 0 e 1.")
    for name in METRICS:
        if name not in EPSILON or not np.isfinite(EPSILON[name]) or EPSILON[name] <= 0:
            raise ValueError(f"EPSILON[{name!r}] deve essere un numero positivo.")
    if any(len(pair) != 2 for pair in IOU_PAIRS):
        raise ValueError("IOU_PAIRS deve essere una matrice N x 2.")
    for first, second in IOU_PAIRS:
        if first not in METRICS or second not in METRICS:
            raise ValueError(f"Coppia IoU ({first}, {second}): metrica non presente in METRICS.")
        if first == second:
            raise ValueError(f"Coppia IoU ({first}, {second}): utilizzare metriche distinte.")
    normalized_pairs = [tuple(sorted(pair)) for pair in IOU_PAIRS]
    if len(normalized_pairs) != len(set(normalized_pairs)):
        raise ValueError("IOU_PAIRS contiene coppie duplicate (anche invertite).")


def main():
    validate_configuration()
    PLOT_DIR.mkdir(parents=True, exist_ok=True)

    values_data = {}
    gradients_data = {}
    gradients_zero_data = {}
    packed_masks = {}
    required_iou_metrics = {name for pair in IOU_PAIRS for name in pair}

    for metric in METRICS:
        npz_path = NPZ_DIR / f"{metric}.npz"
        if not npz_path.is_file():
            raise FileNotFoundError(f"File non trovato: {npz_path}")

        print(f"\nAnalisi di {npz_path.name}...", flush=True)
        with np.load(npz_path, allow_pickle=False) as archive:
            v0 = archive["V0"]
            brt = archive["BRT"]
            coords = tuple(np.asarray(archive[name]) for name in STATE_NAMES)
            periodic_dims = tuple(int(x) for x in np.asarray(archive["periodic_dims"]).flat)
            if v0.ndim != 6 or any(len(c) != n for c, n in zip(coords, v0.shape)):
                raise ValueError(f"{metric}: forma delle coordinate incompatibile con la griglia 6D.")

            values_data[metric] = value_statistics(v0, brt, coords, metric)
            plot_histograms(v0, brt, metric, values_data[metric][12])

            if metric in required_iou_metrics:
                mask = np.isfinite(brt) & (brt <= 0)
                packed_masks[metric] = (
                    np.packbits(mask.ravel()),
                    int(np.count_nonzero(mask)),
                    coords,
                    periodic_dims,
                )

            accuracy = read_solver_accuracy(archive)
            print(f"  Schema derivate: {accuracy}")
            gradients = archive["gradients"]
            gradients_data[metric], gradients_zero_data[metric] = gradient_statistics(
                v0, brt, gradients, coords, periodic_dims, accuracy, EPSILON[metric], metric
            )
            del gradients, v0, brt

    iou_metrics, iou_matrix, iou_pairs = compute_iou(packed_masks, IOU_PAIRS)
    plot_iou_matrix(iou_metrics, iou_matrix)

    workbook = Workbook()
    workbook.active.title = "Value_Function"
    write_sheet(workbook.active, VALUE_ROWS, values_data)
    write_sheet(workbook.create_sheet("Gradients"), GRADIENT_ROWS, gradients_data)
    write_sheet(workbook.create_sheet("Gradients_Zero"), ZERO_ROWS, gradients_zero_data)
    write_iou_sheet(workbook.create_sheet("IoU"), iou_metrics, iou_matrix, iou_pairs)
    workbook.save(OUTPUT_PATH)
    print(f"\nStatistiche salvate in: {OUTPUT_PATH}")
    print(f"Grafici salvati in: {PLOT_DIR}")


if __name__ == "__main__":
    main()
